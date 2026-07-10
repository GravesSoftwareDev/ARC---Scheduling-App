import io
import os
from datetime import date

from openpyxl import load_workbook

from .models import ScheduleEntry

# Microsoft Teams' own "Shifts" import template — starting from it (rather
# than building a workbook from scratch) guarantees the sheet names, header
# row, and column order match exactly what Teams' importer expects.
TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), 'data', 'TeamsShiftsTemplate.xlsx')

# Literal values Teams' importer expects in the "Shared" and "Theme Color"
# columns — these aren't free text, they must match one of Teams' own preset
# option strings exactly.
SHARED_VALUE = '2. Not Shared'
THEME_COLOR = '6. Yellow'


def _fmt_time(t):
    """Format a time as Teams-compatible 12-hour string (e.g. '9am', '1:30pm')."""
    h, m = t.hour, t.minute
    suffix = 'am' if h < 12 else 'pm'
    hour12 = h % 12 or 12
    if m:
        return f"{hour12}:{m:02d}{suffix}"
    return f"{hour12}{suffix}"


def _fmt_date(d):
    """Format a date as M/D/YYYY (no zero-padding, as Teams expects)."""
    return f"{d.month}/{d.day}/{d.year}"


def build_teams_shifts_xlsx(schedule_pks, date_from: date, date_to: date) -> bytes:
    """
    Build a Teams Shifts-compatible xlsx for the given schedules and date range.
    Returns raw bytes ready to serve as a file download.
    """
    wb = load_workbook(TEMPLATE_PATH)

    # Clear sample data from every sheet (keep header row 1)
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.value = None

    shifts_ws = wb['Shifts']

    entries = (
        ScheduleEntry.objects
        .filter(
            schedule_id__in=schedule_pks,
            date__gte=date_from,
            date__lte=date_to,
        )
        .select_related('user', 'schedule')
        .order_by('schedule__name', 'date', 'start_time', 'user__last_name', 'user__first_name')
    )

    # Column order below must match the template's "Shifts" sheet header row
    # exactly: Member, Work Email, Group, Start Date, Start Time, End Date,
    # End Time, Theme Color, Custom Label, Unpaid Break (minutes), Notes, Shared.
    write_row = 2
    for entry in entries:
        emp = entry.user
        lbl = entry.custom_label or None
        if lbl:
            lbl = f'{_fmt_time(entry.start_time)}-{_fmt_time(entry.end_time)} | {entry.custom_label}'
        shifts_ws.cell(row=write_row, column=1).value = emp.get_full_name()
        shifts_ws.cell(row=write_row, column=2).value = emp.email
        shifts_ws.cell(row=write_row, column=3).value = entry.schedule.name
        shifts_ws.cell(row=write_row, column=4).value = _fmt_date(entry.date)
        shifts_ws.cell(row=write_row, column=5).value = _fmt_time(entry.start_time)
        shifts_ws.cell(row=write_row, column=6).value = _fmt_date(entry.date)
        shifts_ws.cell(row=write_row, column=7).value = _fmt_time(entry.end_time)
        shifts_ws.cell(row=write_row, column=8).value = THEME_COLOR
        shifts_ws.cell(row=write_row, column=9).value = lbl
        shifts_ws.cell(row=write_row, column=10).value = None  # Unpaid Break
        shifts_ws.cell(row=write_row, column=11).value = entry.location or None
        shifts_ws.cell(row=write_row, column=12).value = SHARED_VALUE
        write_row += 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
